"""
Stage 5 — Excel Report

Reads raw_results_{flow}.json and scenario_matrix_{flow}.json and writes
a structured .xlsx workbook with three sheets:
  - Scenario Matrix  (all scenarios + expected outcomes)
  - Test Results     (actual outcomes, pass/fail per run)
  - Summary          (category breakdown + overall totals)

Usage:
    python testGen/report.py --flow agent_create_ticket
    python testGen/report.py --flow customer_create_ticket
    python testGen/report.py --flow agent_create_ticket --out path/to/report.xlsx
"""

import argparse
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TESTGEN_DIR = Path(__file__).parent

# ── Colours ───────────────────────────────────────────────────────────────────
C_HEADER    = "1F4E79"   # dark navy
C_PASS      = "E2EFDA"   # light green
C_FAIL      = "FCE4D6"   # light red
C_SKIP      = "FFF2CC"   # light yellow
C_ALT       = "D6E4F0"   # light blue (alternating rows)
C_SUMMARY_H = "2E75B6"   # medium blue
WHITE       = "FFFFFF"

_thin = Side(style="thin", color="BFBFBF")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _hfont(white=True):
    return Font(name="Calibri", bold=True, color=WHITE if white else "000000", size=10)


def _bfont(bold=False, color="000000"):
    return Font(name="Calibri", size=10, bold=bold, color=color)


def _fill(hex_color: str):
    return PatternFill("solid", fgColor=hex_color)


def _col_widths(ws, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header_row(ws, headers: list[str]):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = _hfont()
        cell.fill      = _fill(C_HEADER)
        cell.border    = _border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _inputs_str(inputs: dict | None) -> str:
    if not inputs:
        return ""
    return "\n".join(f"{k}: {v}" for k, v in inputs.items())


def _outcome_fill(val: str) -> PatternFill | None:
    v = str(val).lower()
    if v in ("pass", "submitted", "correctly rejected", "true"):
        return _fill(C_PASS)
    if v in ("fail", "incorrectly accepted", "error", "false"):
        return _fill(C_FAIL)
    if v in ("timeout",):
        return _fill(C_SKIP)
    if v in ("skip", "skipped", "none"):
        return _fill(C_SKIP)
    return None


def _app_behavior_label(actual_outcome: str | None) -> str:
    """Map actual_outcome to a simple PASS / FAIL / SKIP label for the App Behaviour column."""
    v = str(actual_outcome or "").lower().strip()
    if not v or v in ("skipped", "skip", "none", "dry_run"):
        return "SKIP"
    if v in ("submitted", "correctly cascaded", "pass", "incorrectly accepted"):
        return "PASS"
    return "FAIL"


# ── Sheet 1: Scenario Matrix ──────────────────────────────────────────────────

def write_scenario_matrix(wb: Workbook, scenarios: list[dict]):
    ws = wb.create_sheet("Scenario Matrix")
    ws.sheet_view.showGridLines = False

    headers = [
        "ID", "Category", "Flow", "Role", "Test Mode",
        "Description", "Inputs", "Field Under Test",
        "Expected Outcome", "Expected Message", "Business Rule",
    ]
    _col_widths(ws, [8, 20, 22, 12, 14, 42, 32, 18, 16, 42, 12])
    _header_row(ws, headers)

    for r_idx, s in enumerate(scenarios, 2):
        alt = r_idx % 2 == 0
        row = [
            s.get("id", ""),
            s.get("category", ""),
            s.get("flow", ""),
            s.get("role", ""),
            s.get("test_mode", "ui"),
            s.get("description", ""),
            _inputs_str(s.get("inputs")),
            s.get("field") or "",
            s.get("expected_outcome", ""),
            s.get("expected_message") or "",
            s.get("business_rule") or "",
        ]
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font      = _bfont()
            cell.border    = _border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            # Outcome column colour
            if c_idx == 9:
                f = _outcome_fill(str(val))
                cell.fill = f if f else _fill(C_ALT if alt else WHITE)
            else:
                cell.fill = _fill(C_ALT if alt else WHITE)


# ── Sheet 2: Test Results ─────────────────────────────────────────────────────

def write_test_results(wb: Workbook, results: list[dict], run_ts: str) -> int:
    ws = wb.create_sheet("Test Results")
    ws.sheet_view.showGridLines = False

    headers = [
        "ID", "Category", "Role", "Description",
        "Inputs Used", "Expected Scenario", "App Behaviour",
        "Expected Message", "Actual Message",
        "Scenario Result", "Failure Snippet", "Run Timestamp",
    ]
    _col_widths(ws, [8, 20, 10, 42, 32, 16, 20, 42, 42, 16, 60, 20])
    _header_row(ws, headers)

    for r_idx, r in enumerate(results, 2):
        passed = r.get("passed")
        # Scenario Result: did app behaviour match the expected scenario?
        if passed is True:
            sr_label, sr_color, sr_font_color = "✓ PASS", C_PASS, "375623"
        elif passed is False:
            sr_label, sr_color, sr_font_color = "✗ FAIL", C_FAIL, "9C0006"
        else:
            sr_label, sr_color, sr_font_color = "~ SKIP", C_SKIP, "7F6000"

        # App Behaviour: simple PASS / FAIL / SKIP — what the app actually did
        app_beh = _app_behavior_label(r.get("actual_outcome"))
        if app_beh == "PASS":
            ab_color, ab_font_color = C_PASS, "375623"
        elif app_beh == "FAIL":
            ab_color, ab_font_color = C_FAIL, "9C0006"
        else:
            ab_color, ab_font_color = C_SKIP, "7F6000"

        # Extract the most useful snippet from raw pytest output
        raw = r.get("raw_output", "") or ""
        snippet_lines = []
        for line in raw.splitlines():
            stripped = line.strip()
            if stripped.startswith(("FAILED", "E ", "AssertionError", "TimeoutError",
                                    "playwright._impl", "Error:", "assert ")):
                snippet_lines.append(stripped)
            if len(snippet_lines) >= 6:
                break
        failure_snippet = "\n".join(snippet_lines)[:500] if passed is not True else ""

        alt = r_idx % 2 == 0
        row = [
            r.get("id", ""),
            r.get("category", ""),
            r.get("role", ""),
            r.get("description", ""),
            _inputs_str(r.get("inputs")),
            r.get("expected_outcome", ""),
            app_beh,
            r.get("expected_message") or "",
            r.get("actual_message") or "",
            sr_label,
            failure_snippet,
            run_ts,
        ]

        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border    = _border
            cell.alignment = Alignment(wrap_text=True, vertical="top")

            if c_idx == 10:  # Scenario Result
                cell.font  = _bfont(bold=True, color=sr_font_color)
                cell.fill  = _fill(sr_color)
            elif c_idx == 7:  # App Behaviour — PASS / FAIL / SKIP
                cell.font  = _bfont(bold=True, color=ab_font_color)
                cell.fill  = _fill(ab_color)
            elif c_idx == 6:  # Expected Scenario — neutral
                cell.font  = _bfont()
                cell.fill  = _fill(C_ALT if alt else WHITE)
            else:
                cell.font  = _bfont()
                cell.fill  = _fill(C_ALT if alt else WHITE)

    return len(results)


# ── Sheet 3: Summary ──────────────────────────────────────────────────────────

def write_summary(wb: Workbook, results: list[dict], flow: str, run_ts: str):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 10

    def _cell(row, col, value, bold=False, fill_hex=None, center=False, font_color="000000"):
        c = ws.cell(row=row, column=col, value=value)
        c.font      = Font(name="Calibri", size=10, bold=bold, color=font_color)
        c.border    = _border
        c.alignment = Alignment(horizontal="center" if center else "left", vertical="center", wrap_text=True)
        if fill_hex:
            c.fill = _fill(fill_hex)
        return c

    # Title
    ws.merge_cells("A1:E1")
    title_cell = ws.cell(row=1, column=1, value=f"Test Run Summary — {flow}  |  {run_ts}")
    title_cell.font      = Font(name="Calibri", size=12, bold=True, color=WHITE)
    title_cell.fill      = _fill(C_HEADER)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Overall totals block
    total   = len(results)
    passed  = sum(1 for r in results if r.get("passed") is True)
    failed  = sum(1 for r in results if r.get("passed") is False)
    skipped = sum(1 for r in results if r.get("passed") is None)
    pass_rate = f"{passed / total * 100:.1f}%" if total else "—"

    _cell(3, 1, "Overall",       bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE)
    _cell(3, 2, "Total",         bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE, center=True)
    _cell(3, 3, "Passed",        bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE, center=True)
    _cell(3, 4, "Failed",        bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE, center=True)
    _cell(3, 5, "Pass Rate",     bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE, center=True)
    ws.row_dimensions[3].height = 22

    _cell(4, 1, "All scenarios")
    _cell(4, 2, total,     center=True)
    _cell(4, 3, passed,    fill_hex=C_PASS,  center=True, bold=True, font_color="375623")
    _cell(4, 4, failed,    fill_hex=C_FAIL,  center=True, bold=True, font_color="9C0006")
    _cell(4, 5, pass_rate, center=True)

    # Per-category breakdown
    from collections import Counter
    cats = sorted({r.get("category", "") for r in results})
    _cell(7, 1, "By Category",  bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE)
    _cell(7, 2, "Total",        bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE, center=True)
    _cell(7, 3, "Passed",       bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE, center=True)
    _cell(7, 4, "Failed",       bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE, center=True)
    _cell(7, 5, "Skipped",      bold=True, fill_hex=C_SUMMARY_H, font_color=WHITE, center=True)
    ws.row_dimensions[7].height = 22

    for i, cat in enumerate(cats, 8):
        cat_rows = [r for r in results if r.get("category") == cat]
        c_total   = len(cat_rows)
        c_pass    = sum(1 for r in cat_rows if r.get("passed") is True)
        c_fail    = sum(1 for r in cat_rows if r.get("passed") is False)
        c_skip    = sum(1 for r in cat_rows if r.get("passed") is None)
        alt = i % 2 == 0
        bg = C_ALT if alt else WHITE
        _cell(i, 1, cat,     fill_hex=bg)
        _cell(i, 2, c_total, fill_hex=bg, center=True)
        _cell(i, 3, c_pass,  fill_hex=C_PASS if c_pass else bg, center=True, font_color="375623" if c_pass else "000000")
        _cell(i, 4, c_fail,  fill_hex=C_FAIL if c_fail else bg, center=True, font_color="9C0006" if c_fail else "000000")
        _cell(i, 5, c_skip,  fill_hex=C_SKIP if c_skip else bg, center=True, font_color="7F6000" if c_skip else "000000")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Generate Excel report from raw test results")
    parser.add_argument("--flow", required=True, help="Flow name, e.g. agent_create_ticket")
    parser.add_argument("--out",  default=None, help="Output .xlsx path (default: testGen/report_{flow}.xlsx)")
    args = parser.parse_args()

    results_path  = TESTGEN_DIR / f"raw_results_{args.flow}.json"
    matrix_path   = TESTGEN_DIR / f"scenario_matrix_{args.flow}.json"
    out_path      = Path(args.out) if args.out else TESTGEN_DIR / f"report_{args.flow}.xlsx"
    run_ts        = datetime.now().strftime("%Y-%m-%d %H:%M")

    print(f"── Stage 5: Excel Report [{args.flow}] ──────────────────────────────")

    if not results_path.exists():
        print(f"  ERROR: {results_path} not found — run execute.py first.", flush=True)
        return

    results   = json.loads(results_path.read_text(encoding="utf-8"))
    scenarios = []
    if matrix_path.exists():
        raw = json.loads(matrix_path.read_text(encoding="utf-8"))
        scenarios = raw.get("scenarios", raw) if isinstance(raw, dict) else raw

    # Merge scenario role into results (execute.py doesn't persist it)
    role_map = {s.get("id"): s.get("role", "") for s in scenarios}
    for r in results:
        if not r.get("role"):
            r["role"] = role_map.get(r.get("id"), "")

    wb = Workbook()
    wb.remove(wb.active)

    write_summary(wb, results, args.flow, run_ts)
    write_scenario_matrix(wb, scenarios)
    write_test_results(wb, results, run_ts)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))

    passed  = sum(1 for r in results if r.get("passed") is True)
    failed  = sum(1 for r in results if r.get("passed") is False)
    skipped = sum(1 for r in results if r.get("passed") is None)
    total   = len(results)

    print(f"  {passed}/{total} passed  |  {failed} failed  |  {skipped} skipped")
    print(f"  Report saved → {out_path}")


if __name__ == "__main__":
    main()
